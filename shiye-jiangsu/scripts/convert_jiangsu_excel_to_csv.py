"""
将《25年江苏事业单位统考岗位表&竞争比&进面分(1).xlsx》转换为系统可用的标准化CSV。

输入：
- 安徽省事业单位/江苏省事业单位/25年江苏事业单位统考岗位表&竞争比&进面分(1).xlsx

输出：
- 安徽省事业单位/江苏省事业单位/整合后总表_2025年江苏省事业单位岗位.csv

说明：
- 本脚本只负责结构转换与字段清洗，不直接写入数据库。
- 导出后的CSV可直接使用 `scripts/import_positions.py` 进行导入。
"""

import csv
from pathlib import Path
from typing import Optional

import openpyxl


PROJECT_ROOT = Path(__file__).parent.parent

# 源Excel与目标CSV路径
SRC_XLSX = (
    PROJECT_ROOT.parent
    / "安徽省事业单位"
    / "江苏省事业单位"
    / "25年江苏事业单位统考岗位表&竞争比&进面分(1).xlsx"
)

DST_CSV = (
    PROJECT_ROOT.parent
    / "安徽省事业单位"
    / "江苏省事业单位"
    / "整合后总表_2025年江苏省事业单位岗位.csv"
)


FIELDNAMES = [
    "year",
    "exam_type",
    "city",
    "affiliation",
    "region_name",
    "system_type",
    "department_code",
    "department_name",
    "position_code",
    "position_name",
    "position_desc",
    "exam_category",
    "open_ratio",
    "recruit_count",
    "education",
    "major_requirement",
    "other_requirements",
    "apply_count",
    "competition_ratio",
    "min_entry_score",
    "max_entry_score",
]


def _parse_int(value: Optional[object]) -> Optional[int]:
    if value is None:
        return None
    s = str(value).strip()
    if not s:
        return None
    # 去掉千分位等
    s = s.replace(",", "")
    try:
        return int(float(s))
    except ValueError:
        return None


def _parse_float(value: Optional[object]) -> Optional[float]:
    if value is None:
        return None
    s = str(value).strip()
    if not s:
        return None
    s = s.replace(",", "")
    try:
        return float(s)
    except ValueError:
        return None


def _parse_open_ratio(value: Optional[object]) -> Optional[int]:
    """
    将“1:3”之类的开考比例解析为整数（这里取分母，例如 3）。
    """
    if value is None:
        return None
    s = str(value).strip()
    if not s:
        return None
    if ":" in s:
        parts = s.split(":")
        if len(parts) == 2:
            try:
                return int(float(parts[1]))
            except ValueError:
                return None
    # 其他情况尝试直接转为整数
    return _parse_int(s)


def convert():
    if not SRC_XLSX.exists():
        raise FileNotFoundError(f"源Excel不存在: {SRC_XLSX}")

    print("=" * 80)
    print("江苏事业单位岗位表 Excel → 标准化 CSV 转换")
    print("=" * 80)
    print(f"源文件: {SRC_XLSX}")
    print(f"目标文件: {DST_CSV}")

    wb = openpyxl.load_workbook(SRC_XLSX, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]

    # 确保目标目录存在
    DST_CSV.parent.mkdir(parents=True, exist_ok=True)

    with DST_CSV.open("w", encoding="utf-8", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=FIELDNAMES)
        writer.writeheader()

        total = 0
        written = 0

        # 数据从第4行开始（前3行为标题/多行表头）
        for row in ws.iter_rows(min_row=4):
            total += 1
            cells = [cell.value for cell in row]

            # 如果地市、区县、招聘单位名都为空，则认为是空行，跳过
            city = (cells[0] or "").strip() if isinstance(cells[0], str) else cells[0]
            region = (cells[1] or "").strip() if isinstance(cells[1], str) else cells[1]
            unit_name = (
                (cells[3] or "").strip() if isinstance(cells[3], str) else cells[3]
            )

            if not city and not region and not unit_name:
                continue

            # 字段映射（基于表头索引）
            # 索引从0开始：
            # 0 地市, 1 区县, 2 主管部门, 3 招聘单位名称, 4 单位代码, 5 经费来源,
            # 6 岗位名称, 7 岗位代码, 8 笔试类别, 9 岗位描述,
            # 10 招聘人数, 11 开考比例, 12 招聘对象,
            # 13 学历, 14 专业, 15 其他条件,
            # 16 报名人数, 17 竞争比, 18 进面最低分, 19 进面最高分

            supervisor = cells[2]
            department_code = cells[4]
            position_name = cells[6]
            position_code = cells[7]
            exam_category = cells[8]
            position_desc = cells[9]
            recruit_count = _parse_int(cells[10])
            open_ratio = _parse_open_ratio(cells[11])
            recruit_object = cells[12]
            education = cells[13]
            major_requirement = cells[14]
            other = cells[15]
            apply_count = _parse_int(cells[16])
            competition_ratio = _parse_float(cells[17])
            min_entry_score = _parse_float(cells[18])
            max_entry_score = _parse_float(cells[19])

            # 其他条件中附带“招聘对象”信息，避免丢失
            other_text_parts = []
            if recruit_object:
                other_text_parts.append(f"招聘对象：{recruit_object}")
            if other:
                other_text_parts.append(str(other))
            other_combined = "；".join(other_text_parts) if other_text_parts else ""

            record = {
                "year": 2025,
                "exam_type": "事业单位",
                "city": city or None,
                "affiliation": city or None,  # 省属/市属 等，可作为隶属关系
                "region_name": region or None,
                "system_type": supervisor or None,  # 暂用主管部门作为系统类型
                "department_code": str(department_code).strip()
                if department_code not in (None, "")
                else "",
                "department_name": unit_name or "",
                "position_code": str(position_code).strip()
                if position_code not in (None, "")
                else "",
                "position_name": position_name or "",
                "position_desc": position_desc or "",
                "exam_category": exam_category or "",
                "open_ratio": open_ratio if open_ratio is not None else "",
                "recruit_count": recruit_count if recruit_count is not None else 1,
                "education": education or "",
                "major_requirement": major_requirement or "",
                "other_requirements": other_combined,
                "apply_count": apply_count if apply_count is not None else "",
                "competition_ratio": competition_ratio
                if competition_ratio is not None
                else "",
                "min_entry_score": min_entry_score
                if min_entry_score is not None
                else "",
                "max_entry_score": max_entry_score
                if max_entry_score is not None
                else "",
            }

            writer.writerow(record)
            written += 1

            if written % 100 == 0:
                print(f"  已处理 {written} 条记录...")

    print(f"\n✅ 解析行数: {total}")
    print(f"✅ 写入CSV记录数: {written}")
    print("✅ 转换完成，可以使用 scripts/import_positions.py 导入到数据库。")


if __name__ == "__main__":
    convert()

